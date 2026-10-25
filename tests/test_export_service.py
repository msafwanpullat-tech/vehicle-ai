import csv
import json
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook
from PIL import Image as PILImage

from app.routes.counting_routes import download_analysis_export, download_all_analysis_results
from app.services.export_service import generate_asset_zip, generate_exports, generate_results_zip


class ExportServiceTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        self.root = Path(self.temp_dir.name)
        self.analysis_id = "b" * 32
        self.analysis_dir = self.root / "analysis"
        self.analysis_dir.mkdir()
        self.manifest_path = self.root / "session" / "analysis.json"
        self.manifest_path.parent.mkdir()
        self.video_path = self.root / "annotated.mp4"
        self.video_path.write_bytes(b"sample video")
        self.manifest = {
            "analysis_id": self.analysis_id,
            "created_at": "2026-10-21T10:00:00+00:00",
            "video_filename": "traffic.mp4",
            "line_position_percent": 55.6,
            "total_line_crossings": 1,
            "total_tracking_events": 2,
            "class_counts": {"car": 1},
            "anpr_status": "processed",
            "detections": [
                {
                    "timestamp": "00:00:01",
                    "timestamp_seconds": 1.0,
                    "frame_number": 11,
                    "vehicle_id": 9,
                    "vehicle_class": "car",
                    "plate_text": "Uncertain",
                    "ocr_text": "KL07AB1235",
                    "ocr_confidence": 82.4,
                    "status": "Uncertain",
                    "vehicle_image_url": "/outputs/counting/analysis/frame_11_vehicle.jpg",
                    "plate_image_url": "/outputs/counting/analysis/frame_11_plate.jpg",
                }
            ],
            "tracking_events": [
                {
                    "timestamp": "00:00:01",
                    "timestamp_seconds": 1.0,
                    "frame_number": 11,
                    "vehicle_id": 9,
                    "vehicle_class": "car",
                    "line_crossing": True,
                    "crossing_direction": "down",
                    "counted_for_total": True,
                }
            ],
            "search_results": [
                {
                    "query": "KL07AB1234",
                    "message": "No confirmed exact match.",
                    "not_found": False,
                    "exact_matches": [],
                    "uncertain_candidates": [
                        {
                            "vehicle_id": 9,
                            "vehicle_class": "car",
                            "timestamp": "00:00:01",
                            "ocr_text": "KL07AB1235",
                            "ocr_confidence": 82.4,
                            "match_type": "Approximate OCR candidate",
                        }
                    ],
                }
            ],
        }
        self.manifest_path.write_text(json.dumps(self.manifest), encoding="utf-8")
        PILImage.new("RGB", (180, 80), "white").save(
            self.analysis_dir / "frame_11_vehicle.jpg"
        )
        PILImage.new("RGB", (180, 80), "gray").save(
            self.analysis_dir / "frame_11_plate.jpg"
        )

    def tearDown(self) -> None:
        self.temp_dir.cleanup()

    def test_exports_have_expected_csv_and_workbook_contents_and_pdf_header(self) -> None:
        exports = generate_exports(
            self.analysis_id, self.manifest, self.analysis_dir, self.video_path
        )

        with exports["csv"].open(encoding="utf-8-sig", newline="") as stream:
            csv_rows = list(csv.reader(stream))
        self.assertTrue(any(row[0] == "Vehicle Event" for row in csv_rows))
        self.assertTrue(any(row[0] == "Number Plate" for row in csv_rows))
        self.assertTrue(any(row[0] == "Vehicle Search" for row in csv_rows))
        self.assertIn("Uncertain", exports["csv"].read_text(encoding="utf-8-sig"))
        self.assertTrue(all(len(row) == len(csv_rows[0]) for row in csv_rows))
        plate_row = next(row for row in csv_rows if row[0] == "Number Plate")
        self.assertEqual(plate_row[13], "")
        self.assertEqual(plate_row[14], "KL07AB1235")
        class_row = next(row for row in csv_rows if row[0] == "Class Count")
        self.assertEqual(class_row[4], "1")

        workbook = load_workbook(exports["excel"], read_only=False)
        self.assertEqual(
            workbook.sheetnames,
            ["Summary", "Class Counts", "Vehicle Events", "Number Plates", "Vehicle Search Results"],
        )
        self.assertEqual(workbook["Vehicle Events"]["F2"].value, 9)
        self.assertIs(workbook["Vehicle Events"]["J2"].value, True)
        self.assertEqual(workbook["Number Plates"]["J2"].value, "Uncertain")
        self.assertEqual(workbook["Vehicle Search Results"]["B2"].value, "No confirmed exact match.")
        self.assertEqual(workbook["Vehicle Events"].freeze_panes, "A2")
        workbook.close()
        self.assertTrue(exports["pdf"].read_bytes().startswith(b"%PDF"))

    def test_corrupt_screenshot_uses_pdf_placeholder(self) -> None:
        (self.analysis_dir / "frame_11_vehicle.jpg").write_bytes(b"not a valid image")
        exports = generate_exports(
            self.analysis_id, self.manifest, self.analysis_dir, self.video_path
        )
        self.assertTrue(exports["pdf"].read_bytes().startswith(b"%PDF"))

    def test_spreadsheet_export_escapes_formula_like_user_data(self) -> None:
        manifest = {**self.manifest, "video_filename": "=HYPERLINK(\"example\")"}
        exports = generate_exports(
            self.analysis_id, manifest, self.analysis_dir, self.video_path
        )
        workbook = load_workbook(exports["excel"], read_only=False)
        self.assertEqual(
            workbook["Summary"]["B5"].value, "'=HYPERLINK(\"example\")"
        )
        workbook.close()
        with exports["csv"].open(encoding="utf-8-sig", newline="") as stream:
            rows = list(csv.reader(stream))
        video_row = next(row for row in rows if row[3] == "Video filename")
        self.assertEqual(video_row[4], "'=HYPERLINK(\"example\")")

    def test_zip_contains_exports_metadata_images_and_reasonable_video(self) -> None:
        zip_path = self.analysis_dir / "bundle.zip"
        generate_results_zip(
            self.analysis_id,
            self.manifest_path,
            self.manifest,
            self.analysis_dir,
            self.video_path,
            zip_path,
        )

        with zipfile.ZipFile(zip_path) as archive:
            names = set(archive.namelist())
        self.assertTrue({"analysis_results.csv", "analysis_results.xlsx", "analysis_report.pdf"} <= names)
        self.assertIn("analysis.json", names)
        self.assertIn("README.txt", names)
        self.assertIn("vehicle_screenshots/frame_11_vehicle.jpg", names)
        self.assertIn("number_plate_crops/frame_11_plate.jpg", names)
        self.assertIn("annotated_video.mp4", names)

    def test_zip_omits_video_over_100_mib_and_explains_separate_download(self) -> None:
        large_video = self.root / "large.mp4"
        with large_video.open("wb") as stream:
            stream.truncate(100 * 1024 * 1024 + 1)
        zip_path = self.analysis_dir / "large_bundle.zip"
        generate_results_zip(
            self.analysis_id,
            self.manifest_path,
            self.manifest,
            self.analysis_dir,
            large_video,
            zip_path,
        )
        with zipfile.ZipFile(zip_path) as archive:
            self.assertNotIn("annotated_video.mp4", archive.namelist())
            self.assertIn("ANNOTATED_VIDEO.txt", archive.namelist())
            self.assertIn(b"100 MiB", archive.read("ANNOTATED_VIDEO.txt"))

    def test_empty_data_and_missing_assets_do_not_crash(self) -> None:
        empty = {
            "analysis_id": self.analysis_id,
            "video_filename": "empty.mp4",
            "class_counts": {},
            "tracking_events": [],
            "detections": [],
            "search_results": [],
            "total_line_crossings": 0,
            "total_tracking_events": 0,
        }
        exports = generate_exports(
            self.analysis_id, empty, self.analysis_dir, self.video_path
        )
        self.assertGreater(exports["csv"].stat().st_size, 0)
        workbook = load_workbook(exports["excel"], read_only=True)
        self.assertEqual(workbook["Vehicle Events"].max_row, 1)
        self.assertEqual(workbook["Number Plates"].max_row, 1)
        workbook.close()
        asset_zip = self.analysis_dir / "missing_plates.zip"
        generate_asset_zip(self.analysis_dir, "plates", asset_zip)
        with zipfile.ZipFile(asset_zip) as archive:
            self.assertIn("README.txt", archive.namelist())

    def test_download_route_rejects_invalid_format_and_serves_analysis_scoped_file(self) -> None:
        with patch("app.routes.counting_routes._analysis_paths") as path_mock, patch(
            "app.routes.counting_routes._read_export_manifest", return_value=self.manifest
        ), patch("app.routes.counting_routes.generate_exports") as export_mock:
            path_mock.return_value = (
                self.manifest_path,
                self.analysis_dir,
                self.video_path,
            )
            export_mock.return_value = {"csv": self.analysis_dir / "analysis_results.csv"}
            with self.assertRaises(Exception):
                download_analysis_export(self.analysis_id, "unknown")
            export_mock.return_value["csv"].write_text("analysis_id\n", encoding="utf-8")
            response = download_analysis_export(self.analysis_id, "csv")

        self.assertEqual(response.filename, f"vehicle_analysis_{self.analysis_id}.csv")
        self.assertEqual(response.media_type, "text/csv; charset=utf-8")


if __name__ == "__main__":
    unittest.main()
