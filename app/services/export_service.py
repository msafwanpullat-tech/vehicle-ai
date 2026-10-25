import csv
import zipfile
from pathlib import Path
from typing import Any
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from PIL import Image as PILImage
from PIL import UnidentifiedImageError
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import (
    Image,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)
from reportlab.graphics.charts.barcharts import VerticalBarChart
from reportlab.graphics.shapes import Drawing

MAX_VIDEO_IN_BUNDLE_BYTES = 100 * 1024 * 1024
PROJECT_TITLE = (
    "AI-Based Vehicle Detection, Classification and Counting System Using YOLO"
)

SUMMARY_COLUMNS = ["Metric", "Value"]
CLASS_COLUMNS = ["Vehicle Class", "Line-Crossing Count"]
EVENT_COLUMNS = [
    "Analysis ID",
    "Video Filename",
    "Timestamp",
    "Timestamp Seconds",
    "Frame Number",
    "Tracking ID",
    "Vehicle Class",
    "Line Crossing",
    "Crossing Direction",
    "Counted for Total",
]
PLATE_COLUMNS = [
    "Analysis ID",
    "Video Filename",
    "Timestamp",
    "Frame Number",
    "Tracking ID",
    "Vehicle Class",
    "Registration Number",
    "OCR Candidate Text",
    "OCR Confidence (%)",
    "Recognition Status",
    "Vehicle Screenshot",
    "Plate Crop",
]
SEARCH_COLUMNS = [
    "Search Query",
    "Search Status",
    "Match Type",
    "Registration Text",
    "Tracking ID",
    "Vehicle Class",
    "Timestamp",
    "OCR Confidence (%)",
]


def _spreadsheet_safe(value: Any) -> Any:
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def _pdf_text(value: Any) -> str:
    return str(value).encode("cp1252", errors="replace").decode("cp1252")


def _rows(manifest: dict[str, Any], analysis_id: str) -> dict[str, list[list[Any]]]:
    summary = manifest.get("summary", {})
    if not summary:
        summary = {
            key: manifest.get(key, "")
            for key in (
                "analysis_id",
                "created_at",
                "video_filename",
                "line_position_percent",
                "total_line_crossings",
                "total_tracking_events",
                "anpr_status",
                "status",
            )
        }
        summary["class_counts"] = manifest.get("class_counts", {})
    tracking_records = manifest.get("tracking_events")
    tracking_export_status = (
        "Available"
        if isinstance(tracking_records, list)
        else "Not stored for this analysis; no event data was reconstructed."
    )
    summary_pairs: list[list[Any]] = [
        ["Project", PROJECT_TITLE],
        ["Analysis ID", summary.get("analysis_id", analysis_id)],
        ["Analysis date/time (UTC)", summary.get("created_at", "")],
        ["Video filename", summary.get("video_filename", "")],
        ["Analysis status", summary.get("status", "completed")],
        ["Line position (%)", summary.get("line_position_percent", "")],
        ["Total line crossings", summary.get("total_line_crossings", "")],
        ["Tracking observations", summary.get("total_tracking_events", "")],
        ["Tracking event export status", tracking_export_status],
        ["ANPR candidate records", len(manifest.get("detections", []))],
        ["ANPR status", summary.get("anpr_status", "")],
        [
            "Count interpretation",
            "Tracked-ID line-crossing events; not verified unique real-world vehicles.",
        ],
        [
            "Limitations",
            "Occlusion, missed detections, track ID changes, plate visibility and OCR errors may affect results.",
        ],
    ]
    class_rows = [
        [name, count]
        for name, count in sorted(
            (summary.get("class_counts") or {}).items(), key=lambda item: str(item[0]).lower()
        )
    ]
    events = []
    for event in tracking_records or []:
        events.append(
            [
                analysis_id,
                summary.get("video_filename", ""),
                event.get("timestamp", ""),
                event.get("timestamp_seconds", ""),
                event.get("frame_number", ""),
                event.get("vehicle_id", ""),
                event.get("vehicle_class", ""),
                bool(event.get("line_crossing", False)),
                event.get("crossing_direction", ""),
                bool(event.get("counted_for_total", False)),
            ]
        )

    plates = []
    for detection in manifest.get("detections", []):
        plates.append(
            [
                analysis_id,
                summary.get("video_filename", ""),
                detection.get("timestamp", ""),
                detection.get("frame_number", ""),
                detection.get("vehicle_id", ""),
                detection.get("vehicle_class", ""),
                detection.get("plate_text", "")
                if detection.get("status") == "Recognized"
                else "",
                detection.get("ocr_text", ""),
                detection.get("ocr_confidence", ""),
                detection.get("status", "Unreadable"),
                detection.get("vehicle_image_url", ""),
                detection.get("plate_image_url", ""),
            ]
        )

    searches = []
    for search in manifest.get("search_results", []):
        matches = search.get("exact_matches", []) + search.get(
            "uncertain_candidates", []
        )
        if not matches:
            searches.append(
                [
                    search.get("query", ""),
                    "Not found" if search.get("not_found") else "No results",
                    "",
                    "",
                    "",
                    "",
                    "",
                    "",
                ]
            )
        else:
            for match in matches:
                searches.append(
                    [
                        search.get("query", ""),
                        search.get("message", ""),
                        match.get("match_type", ""),
                        match.get("ocr_text", ""),
                        match.get("vehicle_id", ""),
                        match.get("vehicle_class", ""),
                        match.get("timestamp", ""),
                        match.get("ocr_confidence", ""),
                    ]
                )
    return {
        "Summary": summary_pairs,
        "Class Counts": class_rows,
        "Vehicle Events": events,
        "Number Plates": plates,
        "Vehicle Search Results": searches,
    }


def _write_csv(path: Path, data: dict[str, list[list[Any]]], analysis_id: str) -> None:
    columns = [
        "Record Type",
        "Analysis ID",
        "Video Filename",
        "Metric",
        "Value",
        "Vehicle Class",
        "Timestamp",
        "Timestamp Seconds",
        "Frame Number",
        "Tracking ID",
        "Line Crossing",
        "Crossing Direction",
        "Counted for Total",
        "Registration Number",
        "OCR Candidate Text",
        "OCR Confidence (%)",
        "Recognition Status",
        "Search Query",
        "Search Status",
        "Match Type",
        "Vehicle Screenshot",
        "Plate Crop",
    ]
    with path.open("w", newline="", encoding="utf-8-sig") as output:
        writer = csv.writer(output)
        writer.writerow([_spreadsheet_safe(value) for value in columns])
        def write_row(values: list[Any]) -> None:
            writer.writerow([_spreadsheet_safe(value) for value in values])

        summary_rows = data["Summary"]
        filename = next(
            (row[1] for row in summary_rows if row[0] == "Video filename"), ""
        )
        for metric, value in summary_rows:
            write_row(
                ["Summary", analysis_id, filename, metric, value] + [""] * 17
            )
        for vehicle_class, count in data["Class Counts"]:
            write_row(
                ["Class Count", analysis_id, filename, "", count, vehicle_class]
                + [""] * 16
            )
        for event in data["Vehicle Events"]:
            (
                row_analysis_id,
                video_filename,
                timestamp,
                seconds,
                frame,
                track_id,
                vehicle_class,
                crossed,
                direction,
                counted_for_total,
            ) = event
            write_row(
                [
                    "Vehicle Event",
                    row_analysis_id,
                    video_filename,
                    "",
                    "",
                    vehicle_class,
                    timestamp,
                    seconds,
                    frame,
                    track_id,
                    crossed,
                    direction,
                    counted_for_total,
                ]
                + [""] * 9
            )
        for plate in data["Number Plates"]:
            (
                row_analysis_id,
                video_filename,
                timestamp,
                frame,
                track_id,
                vehicle_class,
                registration,
                candidate,
                confidence,
                status,
                vehicle_image,
                plate_image,
            ) = plate
            write_row(
                [
                    "Number Plate",
                    row_analysis_id,
                    video_filename,
                    "",
                    "",
                    vehicle_class,
                    timestamp,
                    "",
                    frame,
                    track_id,
                    "",
                    "",
                    "",
                    registration,
                    candidate,
                    confidence,
                    status,
                    "",
                    "",
                    "",
                    vehicle_image,
                    plate_image,
                ]
            )
        for search in data["Vehicle Search Results"]:
            query, search_status, match_type, text, track_id, vehicle_class, timestamp, confidence = search
            write_row(
                [
                    "Vehicle Search",
                    analysis_id,
                    filename,
                    "",
                    "",
                    vehicle_class,
                    timestamp,
                    "",
                    "",
                    track_id,
                    "",
                    "",
                    "",
                    "",
                    text,
                    confidence,
                    "",
                    query,
                    search_status,
                    match_type,
                    "",
                    "",
                ]
            )


def _write_xlsx(path: Path, data: dict[str, list[list[Any]]]) -> None:
    headers = {
        "Summary": SUMMARY_COLUMNS,
        "Class Counts": CLASS_COLUMNS,
        "Vehicle Events": EVENT_COLUMNS,
        "Number Plates": PLATE_COLUMNS,
        "Vehicle Search Results": SEARCH_COLUMNS,
    }
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, columns in headers.items():
        sheet = workbook.create_sheet(name)
        sheet.append([_spreadsheet_safe(value) for value in columns])
        for row in data[name]:
            sheet.append([_spreadsheet_safe(value) for value in row])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="193B50")
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        for column_cells in sheet.columns:
            width = min(
                54,
                max(
                    12,
                    max(
                        len(str(cell.value or "")) if cell.data_type != "f" else 12
                        for cell in column_cells
                    )
                    + 2,
                ),
            )
            sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = width
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if cell.column in (4, 9) and name == "Vehicle Events":
                    cell.number_format = "0.0"
                if cell.column in (9,) and name == "Number Plates":
                    cell.number_format = "0.0"
        sheet.row_dimensions[1].height = 30
    workbook.save(path)


def _pdf_table(rows: list[list[Any]], widths: list[float] | None = None) -> Table:
    table = Table(rows, colWidths=widths, repeatRows=1, hAlign="LEFT")
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#193B50")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 7),
                ("LEADING", (0, 0), (-1, -1), 9),
                ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#CAD4DE")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F1F5F8")]),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    return table


def _pdf_image_or_placeholder(image_path: Path, styles: Any) -> Any:
    if not image_path.is_file():
        return Paragraph("Image unavailable", styles["BodyText"])
    try:
        with PILImage.open(image_path) as image:
            image.verify()
        image = Image(str(image_path), width=2.7 * inch, height=1.25 * inch, kind="proportional")
        image.wrap(2.7 * inch, 1.25 * inch)
        return image
    except (OSError, UnidentifiedImageError, ValueError):
        return Paragraph("Image unavailable", styles["BodyText"])


def _write_pdf(path: Path, manifest: dict[str, Any], data: dict[str, list[list[Any]]], analysis_dir: Path) -> None:
    styles = getSampleStyleSheet()
    styles.add(
        ParagraphStyle(
            name="ReportTitle",
            parent=styles["Title"],
            alignment=TA_CENTER,
            textColor=colors.HexColor("#193B50"),
            fontSize=17,
        )
    )
    story: list[Any] = [
        Paragraph(PROJECT_TITLE, styles["ReportTitle"]),
        Spacer(1, 12),
    ]
    summary_rows = [["Metric", "Value"]]
    for key, value in data["Summary"]:
        summary_rows.append(
            [
                escape(_pdf_text(key)),
                escape(_pdf_text(value))[:1200],
            ]
        )
    story.extend(
        [
            Paragraph("Analysis Summary", styles["Heading2"]),
            _pdf_table(summary_rows, [2.2 * inch, 5.0 * inch]),
            Spacer(1, 12),
            Paragraph("Vehicle Class Distribution", styles["Heading2"]),
        ]
    )
    class_rows = [["Vehicle class", "Line crossings"]] + [
        [escape(_pdf_text(row[0])), _pdf_text(row[1])] for row in data["Class Counts"]
    ]
    if len(class_rows) == 1:
        class_rows.append(["No class counts available", ""])
    story.extend([_pdf_table(class_rows, [4 * inch, 2 * inch]), Spacer(1, 12)])
    chart_rows = data["Class Counts"]
    if chart_rows:
        chart = VerticalBarChart()
        chart.x = 28
        chart.y = 26
        chart.height = 112
        chart.width = 490
        chart.data = [[int(row[1]) for row in chart_rows]]
        chart.categoryAxis.categoryNames = [_pdf_text(row[0])[:14] for row in chart_rows]
        chart.categoryAxis.labels.fontSize = 6
        chart.categoryAxis.labels.angle = 25
        chart.categoryAxis.labels.boxAnchor = "ne"
        chart.valueAxis.valueMin = 0
        chart.valueAxis.valueMax = max(1, max(int(row[1]) for row in chart_rows))
        chart.valueAxis.valueStep = max(1, chart.valueAxis.valueMax // 5)
        chart.bars[0].fillColor = colors.HexColor("#37A98A")
        chart.barWidth = min(18, 350 / max(1, len(chart_rows)))
        chart.groupSpacing = 4
        drawing = Drawing(540, 165)
        drawing.add(chart)
        story.extend([Paragraph("Class Count Chart", styles["Heading3"]), drawing, Spacer(1, 8)])

    detections = manifest.get("detections", [])
    plate_rows = [[
        "Timestamp",
        "Track ID",
        "Class",
        "Registration/status",
        "Confidence",
    ]]
    for item in detections:
        registration = (
            item.get("plate_text", "")
            if item.get("status") == "Recognized"
            else item.get("status", "Unreadable")
        )
        plate_rows.append(
            [
                escape(_pdf_text(item.get("timestamp", ""))),
                escape(_pdf_text(item.get("vehicle_id", ""))),
                escape(_pdf_text(item.get("vehicle_class", ""))),
                escape(_pdf_text(registration)),
                f"{float(item.get('ocr_confidence', 0)):.1f}%",
            ]
        )
    if len(plate_rows) == 1:
        plate_rows.append(["-", "-", "-", "No plate candidates", "-"])
    story.extend(
        [
            Paragraph("Number-Plate OCR Results", styles["Heading2"]),
            _pdf_table(plate_rows, [0.9 * inch, 0.6 * inch, 1.2 * inch, 2.5 * inch, 1.0 * inch]),
            Spacer(1, 10),
        ]
    )

    events = manifest.get("tracking_events", [])
    event_rows = [
        ["Timestamp", "Track ID", "Class", "Line crossing", "Counted"]
    ]
    for event in events[:500]:
        event_rows.append(
            [
                escape(_pdf_text(event.get("timestamp", ""))),
                escape(_pdf_text(event.get("vehicle_id", ""))),
                escape(_pdf_text(event.get("vehicle_class", ""))),
                escape(
                    _pdf_text(
                        event.get("crossing_direction", "No")
                        or ("Crossing" if event.get("line_crossing") else "No")
                    )
                ),
                "Yes" if event.get("counted_for_total", False) else "No",
            ]
        )
    if len(event_rows) == 1:
        event_rows.append(["-", "-", "No tracking events", "-", "-"])
    story.extend(
        [
            Paragraph("Tracking Observations (first 500)", styles["Heading2"]),
            _pdf_table(
                event_rows,
                [1.2 * inch, 0.8 * inch, 2.2 * inch, 1.5 * inch, 0.7 * inch],
            ),
        ]
    )

    searches = manifest.get("search_results", [])
    if searches:
        search_rows = [
            [
                "Query",
                "Search status",
                "Match type",
                "OCR text",
                "Track ID",
                "Class",
                "Timestamp",
                "Confidence",
            ]
        ]
        for search in searches:
            matches = [
                ("Exact", item) for item in search.get("exact_matches", [])
            ] + [
                ("Uncertain", item)
                for item in search.get("uncertain_candidates", [])
            ]
            if not matches:
                search_rows.append(
                    [
                        escape(_pdf_text(search.get("query", ""))),
                        escape(_pdf_text(search.get("message", "")))[:200],
                        "",
                        "",
                        "",
                        "",
                        "",
                        "",
                    ]
                )
            for match_type, match in matches:
                search_rows.append(
                    [
                        escape(_pdf_text(search.get("query", ""))),
                        escape(_pdf_text(search.get("message", "")))[:200],
                        match_type,
                        escape(_pdf_text(match.get("ocr_text", ""))),
                        escape(_pdf_text(match.get("vehicle_id", ""))),
                        escape(_pdf_text(match.get("vehicle_class", ""))),
                        escape(_pdf_text(match.get("timestamp", ""))),
                        f"{float(match.get('ocr_confidence', 0)):.1f}%",
                    ]
                )
        story.extend(
            [
                PageBreak(),
                Paragraph("Vehicle Search Results", styles["Heading2"]),
                _pdf_table(
                    search_rows,
                    [
                        1.1 * inch,
                        2.2 * inch,
                        0.8 * inch,
                        1.2 * inch,
                        0.5 * inch,
                        0.8 * inch,
                        0.8 * inch,
                        0.6 * inch,
                    ],
                ),
            ]
        )

    image_items = [
        item
        for item in detections
        if item.get("vehicle_image_url") or item.get("plate_image_url")
    ][:3]
    if image_items:
        story.extend([PageBreak(), Paragraph("Selected ANPR Screenshots", styles["Heading2"])])
        for item in image_items:
            story.append(
                Paragraph(
                    f"Track {escape(_pdf_text(item.get('vehicle_id', '')))} - "
                    f"{escape(_pdf_text(item.get('timestamp', '')))} - "
                    f"{escape(_pdf_text(item.get('status', '')))}",
                    styles["BodyText"],
                )
            )
            vehicle_path = analysis_dir / Path(
                str(item.get("vehicle_image_url", "")).rsplit("/", 1)[-1]
            ).name
            plate_path = analysis_dir / Path(
                str(item.get("plate_image_url", "")).rsplit("/", 1)[-1]
            ).name
            image_row = []
            for image_path in (vehicle_path, plate_path):
                image_row.append(_pdf_image_or_placeholder(image_path, styles))
            story.extend([_pdf_table([image_row]), Spacer(1, 8)])

    SimpleDocTemplate(
        str(path),
        pagesize=landscape(letter),
        rightMargin=0.45 * inch,
        leftMargin=0.45 * inch,
        topMargin=0.5 * inch,
        bottomMargin=0.5 * inch,
        title=PROJECT_TITLE,
    ).build(story)


def generate_exports(
    analysis_id: str,
    manifest: dict[str, Any],
    analysis_dir: Path,
    output_video: Path,
) -> dict[str, Path]:
    data = _rows(manifest, analysis_id)
    csv_path = analysis_dir / "analysis_results.csv"
    xlsx_path = analysis_dir / "analysis_results.xlsx"
    pdf_path = analysis_dir / "analysis_report.pdf"
    _write_csv(csv_path, data, analysis_id)
    _write_xlsx(xlsx_path, data)
    _write_pdf(pdf_path, manifest, data, analysis_dir)
    return {"csv": csv_path, "excel": xlsx_path, "pdf": pdf_path}


def generate_asset_zip(analysis_dir: Path, kind: str, destination: Path) -> Path:
    pattern = "*_vehicle.jpg" if kind == "vehicles" else "*_plate.jpg"
    files = sorted(path for path in analysis_dir.glob(pattern) if path.is_file())
    with zipfile.ZipFile(destination, "w", zipfile.ZIP_DEFLATED) as archive:
        for path in files:
            archive.write(path, arcname=path.name)
        archive.writestr(
            "README.txt",
            f"Contains {kind} screenshots for one analysis. Images may be missing if no candidate was recorded.\n",
        )
    return destination


def generate_results_zip(
    analysis_id: str,
    manifest_path: Path,
    manifest: dict[str, Any],
    analysis_dir: Path,
    output_video: Path,
    destination: Path,
) -> Path:
    exports = generate_exports(analysis_id, manifest, analysis_dir, output_video)
    zip_readme = analysis_dir / "EXPORT_README.txt"
    zip_readme.write_text(
        "Vehicle AI analysis export contents:\n"
        "- analysis_results.csv: combined summary, class, tracking, OCR, and search rows\n"
        "- analysis_results.xlsx: separate Summary, Class Counts, Vehicle Events, Number Plates, and Vehicle Search Results sheets\n"
        "- analysis_report.pdf: analysis summary, counts, tracking and OCR results, searches, sample crops, and limitations\n"
        "- *_vehicle.jpg and *_plate.jpg: sampled local ANPR crops\n"
        "- analysis.json: raw per-analysis data used for exports\n"
        "Line counts represent tracked-ID crossing events, not verified unique real-world vehicles. "
        "OCR results marked uncertain are not confirmed registrations. Older analyses may not "
        "contain saved tracking or class-count data; unavailable values are left blank rather "
        "than reconstructed by rerunning video inference.\n",
        encoding="utf-8",
    )
    with zipfile.ZipFile(destination, "w", zipfile.ZIP_DEFLATED) as archive:
        for path in exports.values():
            archive.write(path, arcname=path.name)
        archive.write(manifest_path, arcname="analysis.json")
        archive.write(zip_readme, arcname="README.txt")
        for path in sorted(analysis_dir.glob("*_vehicle.jpg")):
            if path.is_file():
                archive.write(path, arcname=f"vehicle_screenshots/{path.name}")
        for path in sorted(analysis_dir.glob("*_plate.jpg")):
            if path.is_file():
                archive.write(path, arcname=f"number_plate_crops/{path.name}")
        if output_video.is_file():
            if output_video.stat().st_size <= MAX_VIDEO_IN_BUNDLE_BYTES:
                archive.write(output_video, arcname="annotated_video" + output_video.suffix)
            else:
                archive.writestr(
                    "ANNOTATED_VIDEO.txt",
                    "The annotated video exceeded the ZIP size threshold (100 MiB). Download it separately from the dashboard.\n",
                )
    return destination
